from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from ocr_translation.errors import DependencyMissingError
from ocr_translation.models import TextBlock


@dataclass
class DocxWriter:
    title: str

    def write_blocks(self, output_path: Path, blocks: list[TextBlock]) -> None:
        try:
            from docx import Document
        except ImportError as exc:
            raise DependencyMissingError(
                "python-docx is not installed. Install requirements.txt before writing DOCX files."
            ) from exc

        document = Document()
        document.add_heading(self.title, level=1)

        current_page: int | None = None
        for block in blocks:
            if block.page_number != current_page:
                current_page = block.page_number
                document.add_heading(f"Page {current_page}", level=2)

            for paragraph in _split_paragraphs(block.text):
                document.add_paragraph(paragraph)

        document.save(output_path)

    def write_paragraphs(self, output_path: Path, paragraphs: list[str]) -> None:
        try:
            from docx import Document
        except ImportError as exc:
            raise DependencyMissingError(
                "python-docx is not installed. Install requirements.txt before writing DOCX files."
            ) from exc

        document = Document()
        document.add_heading(self.title, level=1)
        for paragraph in paragraphs:
            if paragraph.strip():
                document.add_paragraph(paragraph.strip())
        document.save(output_path)


def _split_paragraphs(text: str) -> list[str]:
    paragraphs = [part.strip() for part in text.splitlines() if part.strip()]
    return paragraphs or [text.strip()]
